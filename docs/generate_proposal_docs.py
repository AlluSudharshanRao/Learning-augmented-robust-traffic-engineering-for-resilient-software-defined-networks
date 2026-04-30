from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer


TITLE = "Learning-Augmented Robust Traffic Engineering for Resilient Software-Defined Networks"
SUBTITLE = "Final Project Proposal for ECE GY 7363 Communications Networks II: Design and Algorithms"

SECTIONS = [
    (
        "1. Problem Definition",
        [
            "Modern communication networks face two major challenges: traffic demand changes over time and link or node failures can suddenly degrade performance.",
            "Traditional routing methods such as shortest path routing are simple, but they do not adapt well to dynamic traffic conditions. Pure optimization-based routing is stronger, but if it only reacts to current traffic, it may still perform poorly when demand changes rapidly.",
            "This project proposes a learning-augmented network design framework that combines machine learning for traffic matrix prediction, multi-commodity flow optimization for routing, and robust design principles for failure-aware performance.",
            "The main question is whether traffic prediction can improve routing quality and network resilience compared with classical routing and non-predictive optimization methods.",
        ],
    ),
    (
        "2. Relevance to Course Topics",
        [
            "This project directly matches the course themes: network design problem modeling, optimization methods, multi-commodity flow routing, fair network design, resilient network design, robust network design, and machine-learning-based network management.",
            "It also follows the final project format requiring a clear problem definition, design models, solution techniques, and performance evaluation.",
        ],
    ),
    (
        "3. Project Objectives",
        [
            "Build a time-varying traffic engineering problem on a communication network.",
            "Train an ML model to predict short-term future traffic demands.",
            "Use predicted demand in an optimization model for routing.",
            "Incorporate robustness against link failures.",
            "Compare the proposed method with classical and optimization-only baselines.",
        ],
    ),
    (
        "4. Proposed System",
        [
            "At each time step, the system will observe recent traffic matrices, predict the next traffic matrix, feed that prediction into a routing optimizer, compute feasible flow allocations over the network, and evaluate performance under normal operation and failure scenarios.",
            "The end-to-end pipeline is: historical traffic -> ML prediction -> optimization-based routing -> resilience evaluation.",
        ],
    ),
    (
        "5. Mathematical Design Model",
        [
            "We model the network as a directed graph G = (V, E), where nodes represent routers or switches and links have capacities.",
            "Traffic demands are represented as commodities, one for each source-destination pair.",
            "Decision variables represent the amount of each commodity routed on each link.",
            "The optimization model will enforce flow conservation, link capacity limits, and nonnegativity constraints.",
            "The primary objective is to minimize maximum link utilization, with optional alternatives such as minimizing total congestion cost or delay.",
        ],
    ),
    (
        "6. Robust and Resilient Extension",
        [
            "To make the design failure-aware, we will evaluate routing under link-failure scenarios such as random single-link failures and critical-link failures.",
            "The baseline implementation will optimize routing for the nominal network and then perform scenario-based resilience evaluation.",
            "If time permits, we will extend the model to a simplified robust optimization formulation that accounts for a predefined set of failure scenarios directly in the solver.",
        ],
    ),
    (
        "7. Machine Learning Component",
        [
            "The ML module predicts the next traffic matrix using historical traffic observations.",
            "We will compare simple predictors such as moving average or linear regression against a sequence model such as LSTM or GRU.",
            "The recommended main model is an LSTM because it captures temporal traffic patterns while remaining implementable within a semester project.",
        ],
    ),
    (
        "8. Algorithms and Baselines",
        [
            "Baseline 1: shortest path routing with no optimization or prediction.",
            "Baseline 2: multi-commodity flow optimization using only current observed demand.",
            "Baseline 3: robust or conservative optimization without an ML predictor.",
            "Proposed method: learning-augmented routing that uses predicted demand inside the optimization model.",
        ],
    ),
    (
        "9. Fairness Component",
        [
            "To connect the project with fair network design, we will include fairness in evaluation through metrics such as Jain's fairness index.",
            "If time allows, we will also add a fairness-aware objective or constraint and compare the tradeoff between throughput, congestion, and fairness.",
        ],
    ),
    (
        "10. Data and Network Instances",
        [
            "We will evaluate on standard communication-network topologies such as NSFNET, Abilene, GEANT, or a synthetic stress-test topology.",
            "Because public time-varying traffic traces may be limited, we will generate realistic traffic matrices using periodic demand trends, random noise, bursty traffic spikes, and hotspot source-destination pairs.",
            "The generated data will be split into training data for ML, test data for routing evaluation, and stress scenarios for resilience analysis.",
        ],
    ),
    (
        "11. Performance Metrics",
        [
            "Maximum link utilization",
            "Average link utilization",
            "Fraction of demand satisfied",
            "Total routing cost",
            "Average path length",
            "Jain's fairness index",
            "Number of congested links",
            "Performance degradation under failure",
            "Prediction error of the ML model",
        ],
    ),
    (
        "12. Experimental Scenarios",
        [
            "Normal operation with moderate traffic variation",
            "Traffic spike scenarios with hotspot demands",
            "Random and critical single-link failures",
            "Robustness tests under prediction error and noisy traffic patterns",
        ],
    ),
    (
        "13. Expected Contributions",
        [
            "A complete end-to-end learning-plus-optimization framework for traffic engineering.",
            "A rigorous comparison against shortest path, optimization-only, and robust non-predictive baselines.",
            "A resilience analysis showing when prediction helps and when conservative optimization remains preferable.",
        ],
    ),
    (
        "14. Tools and Implementation Plan",
        [
            "Programming language: Python",
            "Graph handling: networkx",
            "Data processing: numpy and pandas",
            "Optimization: cvxpy or pulp, with Gurobi or CPLEX if available",
            "Machine learning: PyTorch",
            "Visualization: matplotlib or seaborn",
        ],
    ),
    (
        "15. Step-by-Step Work Plan",
        [
            "Phase 1: select topology, define traffic model, and implement shortest-path baseline.",
            "Phase 2: formulate and validate the multi-commodity flow optimization model.",
            "Phase 3: generate time-series traffic matrices and train baseline and LSTM predictors.",
            "Phase 4: integrate predicted traffic into the optimizer and compare with non-predictive routing.",
            "Phase 5: simulate failures, evaluate resilience, and add the robust design discussion or extension.",
            "Phase 6: prepare final report, plots, slides, and clean submission package.",
        ],
    ),
    (
        "16. Week-by-Week Timeline",
        [
            "Week 1: finalize problem statement, choose topology, and set up the codebase.",
            "Week 2: implement shortest-path routing and synthetic traffic generation.",
            "Week 3: build the multi-commodity flow optimization model.",
            "Week 4: train baseline and LSTM traffic predictors.",
            "Week 5: integrate prediction with routing optimization.",
            "Week 6: run failure and robustness experiments.",
            "Week 7: analyze results and prepare report and slides.",
        ],
    ),
    (
        "17. Deliverables",
        [
            "Project report with problem definition, mathematical model, algorithms, experiments, and lessons learned.",
            "Presentation slides covering motivation, formulation, baselines, and results.",
            "Code with a clear folder structure, run scripts, and example datasets.",
            "Traffic matrices, failure scenarios, and topology descriptions used in experiments.",
        ],
    ),
    (
        "18. Risk Management",
        [
            "Optimization may become slow for large networks, so we will begin with medium-size topologies and scale carefully.",
            "ML may not help if traffic is too random, so we will keep strong baselines and analyze both successes and failures honestly.",
            "Robust optimization may be too complex to fully implement, so scenario-based resilience evaluation will be completed first.",
        ],
    ),
    (
        "19. Why This Project Can Score High",
        [
            "Novelty and relevance come from combining ML, traffic engineering, and resilience in one coherent framework.",
            "Depth and workload come from building modeling, optimization, ML, and simulation components together.",
            "Correctness is supported by a mathematically grounded optimization core and systematic evaluation.",
            "Presentation quality will be strong because the work naturally produces clear plots, topology figures, and baseline comparisons.",
        ],
    ),
    (
        "20. Proposed Abstract",
        [
            "This project studies learning-augmented traffic engineering for resilient software-defined networks. We consider a network with time-varying traffic demands and possible link failures. A machine learning model predicts the next traffic matrix using historical traffic observations, and the predicted demand is then used in a multi-commodity flow optimization model to compute routing decisions that minimize network congestion. We compare the proposed framework against shortest-path routing, optimization using only current demand, and robust non-predictive baselines. Performance is evaluated on standard network topologies under dynamic traffic and link-failure scenarios using metrics such as maximum link utilization, demand satisfaction, fairness, and resilience. The goal is to determine whether prediction-enhanced optimization can improve routing efficiency and robustness in practical network design settings.",
        ],
    ),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.15):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_bullets_docx(document, items):
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        run.font.size = Pt(10.5)
        set_paragraph_spacing(p, after=2, line=1.15)


def build_docx(path: Path):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title, after=8, line=1.0)
    r = title.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(23, 54, 93)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(subtitle, after=10, line=1.0)
    r = subtitle.add_run(SUBTITLE)
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(90, 90, 90)

    summary = doc.add_table(rows=2, cols=2)
    summary.style = "Table Grid"
    summary.autofit = False
    summary.columns[0].width = Inches(2.0)
    summary.columns[1].width = Inches(4.8)
    entries = [
        ("Project Type", "Machine learning plus optimization for resilient network design"),
        ("Core Deliverable", "Traffic prediction model integrated with multi-commodity flow routing and failure-aware evaluation"),
    ]
    for row, (label, value) in zip(summary.rows, entries):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], "D9EAF7")
        for idx, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                set_paragraph_spacing(p, after=0, line=1.1)
                for run in p.runs:
                    run.font.size = Pt(10)
                    if idx == 0:
                        run.bold = True

    doc.add_paragraph("")

    for heading, bullets in SECTIONS:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=4, after=2, line=1.0)
        run = p.add_run(heading)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(23, 54, 93)
        add_bullets_docx(doc, bullets)

    doc.save(path)


def build_pdf(path: Path):
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleCustom",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#17365D"),
        alignment=1,
        spaceAfter=8,
    )
    subtitle_style = ParagraphStyle(
        "SubtitleCustom",
        parent=styles["BodyText"],
        fontName="Helvetica-Oblique",
        fontSize=10.5,
        leading=14,
        textColor=colors.HexColor("#555555"),
        alignment=1,
        spaceAfter=12,
    )
    heading_style = ParagraphStyle(
        "HeadingCustom",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=15,
        textColor=colors.HexColor("#17365D"),
        spaceBefore=8,
        spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "BodyCustom",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10,
        leading=13,
        spaceAfter=2,
    )

    story = [
        Paragraph(TITLE, title_style),
        Paragraph(SUBTITLE, subtitle_style),
        Paragraph("<b>Project Type:</b> Machine learning plus optimization for resilient network design", body_style),
        Paragraph("<b>Core Deliverable:</b> Traffic prediction model integrated with multi-commodity flow routing and failure-aware evaluation", body_style),
        Spacer(1, 0.12 * inch),
    ]

    for heading, bullets in SECTIONS:
        story.append(Paragraph(heading, heading_style))
        list_items = [ListItem(Paragraph(item, body_style), leftIndent=10) for item in bullets]
        story.append(
            ListFlowable(
                list_items,
                bulletType="bullet",
                start="circle",
                leftIndent=16,
            )
        )
        story.append(Spacer(1, 0.06 * inch))

    doc = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        leftMargin=0.8 * inch,
        rightMargin=0.8 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch,
    )
    doc.build(story)


def main():
    out_dir = Path(__file__).resolve().parent
    build_docx(out_dir / "ML_Network_Project_Proposal.docx")
    build_pdf(out_dir / "ML_Network_Project_Proposal.pdf")


if __name__ == "__main__":
    main()
